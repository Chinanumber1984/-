from pathlib import Path
import csv
root=Path(__file__).resolve().parents[1];base=root/'docs/review/manual-acceptance/W20-22-质量风险验收用例';base.parent.mkdir(parents=True,exist_ok=True)
rows=[
('P0','lixiaolong','PPAP首版提交','质量与风险→新建PPAP，填标题；上传PDF／PNG／JPEG，填写每份提交物名称和说明→提交','状态待SQE审核；无附件或空清单应被拒绝'),
('P0','lijunhua','SQE要求补件','SQE待办打开上一份PPAP，填写审核意见→补件／重做','门户看到补件意见；第一版记录保留'),
('P0','lixiaolong / admin','补件重提并批准','门户上传补充报告并重提；admin在SQE待办审核通过','PPAP已批准，版本2；历史第一版补件状态和附件保留'),
('P0','gysadmin','门户管理员查看质量资料','打开同一PPAP并下载附件','可见本供方资料并成功下载；无内部风险、SQE待办入口'),
('P0','lijunhua','发起SCAR','选择当前门户绑定供方，填写独立测试标题、缺陷编码、事实和严重度','状态待8D回复，显示按创建时规则计算的回复截止'),
('P0','lixiaolong','8D必填与附件','打开本方SCAR；先留空一个D步骤尝试提交，再补齐D1–D8及附件后提交','缺步或无附件时拦截；完整结构化回复进入SQE审核'),
('P0','lijunhua / lixiaolong','8D重做','SQE补件／重做并说明；门户修正根因和防再发内容再提交','旧版保留，新版独立审核；不能覆盖原回复'),
('P0','admin / lijunhua','验证与有效性闭环','admin审核8D通过；lijunhua填写措施有效性结果和证据引用→闭环','先已验证、待有效性确认，再已闭环；未验证的SCAR不能直接关闭'),
('P1','lijunhua','重复缺陷升级','对同供方发起两张相同缺陷编码的独立测试SCAR；等待约30秒','生成重复缺陷内部风险及双方质量通知；不直接冻结'),
('P1','admin','响应期限规则版本','质量规则发布新响应天数及重复窗口，写调整依据；再建SCAR','新单用新规则；已有SCAR保留创建时的期限；超期未回复将升级'),
('P1','lijunhua','内部风险归集','查看内部风险，核对之前采购协同的超容差／逾期事件，以及过期资质事件','来源固定INTERNAL，有原单据或资质引用；待评估不等于低风险'),
('P0','lijunhua','有据人工合规登记','新建明显标为“验收演示，非真实黑名单”的风险，填事实和证据引用','无证据不能保存；来源为INTERNAL，保留登记记录'),
('P0','lijunhua','等级评估与红黄灯','上述风险评估为高，说明依据；再检查内部风险列表和原寻源风险摘要','显示红灯和经理冻结建议，寻源风险摘要反映当前事件；供方未自动冻结'),
('P0','lijunhua / admin','冻结建议接原流程','只对新建并已准入的隔离供应商登记高风险；提交冻结建议后，在原审批工作台审批','提交时仍合格；本人不能审批本人申请；admin批准后才冻结，后续解冻也须原审批'),
('P0','admin','风险处置归档','对验收演示风险填写处置意见和证据，选择整改完成或有据观察并归档','处置留痕保留；高风险须经理角色归档，不自动解除已批准冻结'),
('P0','lijunhua','来源缺失不补分','检验与绩效来源，选择新建隔离供应商和空月份→预览','提示需要ERP基础数据及QMS检验，不用满分／零分代替缺失数据'),
('P0','lijunhua','沙箱ERP与QMS联调','对有映射的隔离供应商，模拟ERP基础项、QMS严重不合格；等待几秒后预览','显示本机沙箱；默认严重扣20，质量80；检验引用、来源版本与规则版本可追溯'),
('P0','lijunhua','幂等评分与主档隔离','记录隔离供方原等级；按当前来源发布并重算，再点一次；检查原绩效及供方等级','相同依据不新增重复输入；有SANDBOX_ERP_QMS标签；沙箱评分不会联动主档等级'),
('P0','lijunhua','旧依据拦截','预览之后再模拟一条QMS不合格，不刷新预览直接发布；然后重新预览','旧依据发布拦截；新预览扣分随当前有效检验变化，最低0分'),
('P1','lijunhua','真实／沙箱隔离','来源模式切为真实接口，查询同供方同月份','不会读取沙箱数据冒充真实分数；真实来源未接入应提示缺失'),
('P0','四个账号','旧功能及权限回归','分别登录检查原资质、寻源、采购协同；门户只看本供方；admin查看集成监控','原数据和密码保留；QMS入站／出站有版本、重试和结果记录，不显示连接密钥'),
('P1','接口负责人','真实厂商联调（待资料）','提供ERP／QMS字段字典、认证方式、样例、测试地址和回调契约后，按真实来源复验','本机沙箱不作为真实联调验收，未提供资料时标记阻塞而非通过')]
headers=['编号','优先级','账号','测试项','步骤','预期结果','实际结果','状态','问题备注'];data=[[f'P7-{i:02d}',*r,'','未测',''] for i,r in enumerate(rows,1)]
for ext in ['.xlsx','.csv','.md']:
 if base.with_suffix(ext).exists():raise SystemExit('Existing quality acceptance results retained; no overwrite')
with base.with_suffix('.csv').open('w',encoding='utf-8-sig',newline='') as f:w=csv.writer(f,lineterminator="\n");w.writerow(headers);w.writerows(data)
intro='''# W20–22 质量风险手工验收\n\n内部：http://localhost:18080/?internal=1\n\n门户：http://localhost:18080/?portal=1\n\n账号和密码沿用当前四个测试账号。请用有明显验收标识的新单据；冻结及评分测试使用新建隔离供方，不改原历史资料。只有当前门户绑定供方可做门户PPAP／8D；QMS／ERP沙箱映射可由管理员配置。\n\n默认规则：8D响应14天，重复缺陷窗口90天；轻微／主要／严重不合格分别扣5／10／20。质量分=max(0,100−当前检验扣分合计)。沙箱结果有明确来源标识，且不联动供应商主档等级。\n\n所有用例初始为未测，请自行填写结果；最后一项真实联调尚待资料，不代表已验收。\n\n'''
with base.with_suffix('.md').open('w',encoding='utf-8') as f:
 f.write(intro)
 for r in data:f.write(f'## {r[0]} {r[3]}（{r[1]}）\n\n账号：{r[2]}\n\n步骤：{r[4]}\n\n预期：{r[5]}\n\n状态：未测\n\n')
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
wb=openpyxl.Workbook();ws=wb.active;ws.title='质量风险验收';ws.append(headers)
for r in data:ws.append(r)
for c in ws[1]:c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='215C47')
for i,width in enumerate([14,10,28,28,78,78,45,14,40],1):ws.column_dimensions[get_column_letter(i)].width=width
for row in ws.iter_rows(min_row=2):
 for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
for i in range(2,len(data)+2):ws.row_dimensions[i].height=95
ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;dv=DataValidation(type='list',formula1='"未测,通过,不通过,阻塞"');ws.add_data_validation(dv);dv.add(f'H2:H{len(data)+1}')
note=wb.create_sheet('说明');note.append([intro]);note.column_dimensions['A'].width=115;note['A1'].alignment=Alignment(wrap_text=True,vertical='top');note.row_dimensions[1].height=330;wb.save(base.with_suffix('.xlsx'));print(f'Generated {len(data)} P7 manual cases; user results and earlier workbooks untouched.')
