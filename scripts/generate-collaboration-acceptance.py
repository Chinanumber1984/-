from pathlib import Path
import csv,io
root=Path(__file__).resolve().parents[1];out=root/'docs/review/manual-acceptance';out.mkdir(parents=True,exist_ok=True)
base=out/'W16-19-采购协同验收用例'
rows=[
('01','P0','lijunhua','生成沙箱订单','采购协同→采购订单→生成本机沙箱订单；选择当前测试供方，数量10、单价100、交期7天后；生成并刷新','显示“本机沙箱”、待供方确认；内部与供方都能找到同一订单'),
('02','P0','lixiaolong','确认当前订单','门户→采购协同→查看本次新订单→确认当前订单','已确认版本等于来源版本；可以填写发运'),
('03','P0','lixiaolong / lijunhua','交期数量异议','供方建议数量12、交期调整并填原因；内部对同一沙箱订单点击“沙箱 ERP 接受建议”并刷新','待ERP处理时不能发运；回写后来源版本+1，供方须重新确认'),
('04','P1','lixiaolong / lijunhua','拒绝订单建议','另建订单，供方建议变更；内部选择“沙箱 ERP 拒绝建议”','保留原有效订单，建议状态已退回，可重新核对履约'),
('05','P0','lixiaolong','两批次发运','确认最新订单后，填写发运通知；同一物料分批次A数量5和B数量7，填写承运商、运单、ETA','保存成功；发运量12；批次和物流可追踪'),
('06','P0','lixiaolong','超量发运拦截','已发满12后尝试再发1，或在新订单上填写超订单量','服务器拒绝；累计已发数量不变'),
('07','P1','lixiaolong','补充物流','已保存且未最终收货的ASN→查看发运／收货→更新物流','内部查看同一ASN可见新承运商、运单和预计到货'),
('08','P0','lijunhua','分次实收','ASN登记人工实收，仅收第一批5，取消“本次为最终收货”','状态部分收货，未收完本身不触发短缺差异'),
('09','P0','lijunhua','最终实收与超差','继续收第二批7.5并勾选最终收货（累计12.5）','状态已收货；超12的3%容差，显示存在差异，双方协同通知有预警'),
('10','P1','admin / lijunhua','容差版本','先建ASN，再把容差改为5%，填写原因；再建新ASN','原ASN仍用创建时容差及版本；后建ASN使用5%；测试后按需恢复3%'),
('11','P1','lijunhua','ERP实收','另建沙箱订单、供方确认并发运；内部点击“模拟 ERP 全量实收”','签名实收回写后显示已收货；重复刷新不会新增实收'),
('12','P1','lijunhua / admin','人工实收冲销','在未对账的人工实收明细填写冲销依据→申请冲销；admin在经理待办批准','申请人不能批准本人；批准后保留冲销记录，累计实收重算；ERP实收不允许人工冲销'),
('13','P0','lijunhua','生成对账','实收与建账勾选本次订单未占用明细，填写可辨识对账编号并生成','数量、金额基于有效实收快照；占用行不可再建另一张账单'),
('14','P0','lijunhua / lixiaolong','双方确认','内部采购确认；供应商财务确认同一账单','只确认一方时仍待双方确认；两方确认后已确认，允许登记发票'),
('15','P0','lijunhua / lixiaolong','差异讨论','在尚未财务受理的账单提出差异并说明数量／金额原因；双方讨论','状态存在差异；双方确认标记清除；讨论意见留痕'),
('16','P0','lijunhua / admin','经理争议审批','差异账单升级经理并填依据；lijunhua尝试批准；admin核查后批准','本人审批被拒；admin批准后账单重新待双方确认，不自动生效'),
('17','P0','lixiaolong','上传发票并校验金额','账单双方确认后上传PDF／PNG／JPEG；登记时先填错误金额，再改为准确金额','错误金额拦截；准确金额及必传附件登记成功；无登录不能下载附件'),
('18','P0','lijunhua','推送共享财务','打开已登记发票→推送共享财务；数秒后刷新协同','先待财务回执，再显示财务已受理；“接口已送达”单独显示，不代表已付款'),
('19','P0','lijunhua','付款申请与回执','已受理发票→提交付款申请；申请送达后，填写一半金额→本机沙箱模拟付款回执；再模拟剩余金额','申请不会直接改为已付款；第一笔部分付款，第二笔已付款；不得超过发票总额'),
('20','P0','lijunhua / lixiaolong','财务受理后保护','已受理或已付款账单再次提出差异、撤回或冲销关联实收','服务器拒绝破坏已进入财务受理的账单；历史金额不变'),
('21','P1','lijunhua','MES与看板','沙箱订单点击“模拟 MES 紧急优先级”；刷新交付看板','可见MES优先级；在途、实收和逾期按有效单据汇总'),
('22','P1','lijunhua / lixiaolong','逾期通知','生成交期为昨天且未收完的沙箱订单，等待约30秒，查看双方协同通知','出现逾期未收预警；有风险待评估记录，不自动冻结供应商'),
('23','P0','gysadmin','门户管理员范围','门户管理员登录→查看本供方订单、对账、发票；检查内部功能入口','可看同供方业务；无内部集成监控和财务推送／付款按钮'),
('24','P1','admin','集成监控','内部管理员→采购协同→集成监控；查看入站、出站和失败重试','只显示事件、版本、状态和错误代码，不显示密钥及报文金额；冲突须来源纠正'),
('25','P0','四个账号','旧功能与数据保留','四个账号分别登录；抽查原供应商、资质附件、原“买点小东西”寻源与历史绩效','原业务保留，密码不变；仍仅四个测试账号启用')]
headers=['编号','优先级','执行账号','测试项','操作步骤','预期结果','实际结果','状态','问题备注']
data=[[f'P6-{n}',prio,who,name,steps,expected,'','未测',''] for n,prio,who,name,steps,expected in rows]
for suffix in ['.csv','.md','.xlsx']:
 if base.with_suffix(suffix).exists():raise SystemExit('Existing manual acceptance file retained; do not overwrite user results')
with base.with_suffix('.csv').open('w',encoding='utf-8-sig',newline='') as f:w=csv.writer(f,lineterminator="\n");w.writerow(headers);w.writerows(data)
intro='''# W16–19 采购协同手工验收\n\n内部入口：http://localhost:18080/?internal=1\n\n供应商入口：http://localhost:18080/?portal=1\n\n四个账号及当前密码沿用现有设置。请使用本次新建且有明显标识的沙箱订单，不修改原有历史订单。建议按用例顺序操作，先完成 P0；实际结果与状态由验收人员填写。\n\n本机沙箱会生成真实持久化的 SRM 业务记录，但不代表已对接实际 ERP、MES 或财务共享，更不会实际付款。真实外部系统联调仍待接口资料；风险事件的评估消费属于后续风险阶段。\n\n'''
with base.with_suffix('.md').open('w',encoding='utf-8') as f:
 f.write(intro)
 for row in data:f.write(f'## {row[0]} {row[3]}（{row[1]}）\n\n账号：{row[2]}\n\n步骤：{row[4]}\n\n预期：{row[5]}\n\n结果：未测\n\n')
try:
 import openpyxl
 from openpyxl.styles import Font,PatternFill,Alignment
 from openpyxl.worksheet.datavalidation import DataValidation
 from openpyxl.utils import get_column_letter
 wb=openpyxl.Workbook();ws=wb.active;ws.title='采购协同验收';ws.append(headers)
 for row in data:ws.append(row)
 for c in ws[1]:c.fill=PatternFill('solid',fgColor='215C47');c.font=Font(color='FFFFFF',bold=True)
 widths=[14,10,28,24,72,72,45,14,40]
 for i,w in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=w
 for row in ws.iter_rows(min_row=2):
  for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
 for i in range(2,len(data)+2):ws.row_dimensions[i].height=90
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 dv=DataValidation(type='list',formula1='"未测,通过,不通过,阻塞"');ws.add_data_validation(dv);dv.add(f'H2:H{len(data)+1}')
 note=wb.create_sheet('说明');note.append(['W16–19 采购协同验收说明']);note.append([intro]);note.column_dimensions['A'].width=110;note['A2'].alignment=Alignment(wrap_text=True,vertical='top');note.row_dimensions[2].height=260;wb.save(base.with_suffix('.xlsx'))
except ImportError:print('CSV and Markdown generated; openpyxl unavailable, XLSX not generated.')
print(f'Generated {len(data)} manual cases; all initially 未测. Existing W12–15 workbook untouched.')
