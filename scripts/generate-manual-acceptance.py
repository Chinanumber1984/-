"""Generate a fillable local acceptance workbook. Does not access accounts or the database.
Dependency: openpyxl==3.1.5. Refuses to overwrite files that may contain user results.
"""
from pathlib import Path
from datetime import datetime
import csv,secrets,sys,zipfile
from xml.etree import ElementTree as ET
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

out=Path(sys.argv[1]) if len(sys.argv)>1 else Path(__file__).resolve().parents[1]/'docs/review/manual-acceptance'
base='W12-15-手工验收用例'
paths=[out/(base+ext) for ext in ['.xlsx','.md','.csv']]+[out/'W12-15-验收规格示例.pdf']
if any(p.exists() for p in paths):raise SystemExit('目标文件已存在；请另选输出目录，避免覆盖已填写的验收结果。')
out.mkdir(parents=True,exist_ok=True)
headers=['用例编号','优先级','场景','账号','前置条件','操作步骤','测试数据','预期结果','关联单号','实际结果','测试结论','截图/备注']
cases=[]
def case(cid,priority,title,actor,pre,steps,data,expected):
    cases.append([cid,priority,title,actor,pre,steps,data,expected,'','','未测',''])
P0='P0 必测';P1='P1 选测'
case('M01',P0,'四账号登录与企业核对','四个账号分别登录','使用说明中的两个入口；验证码照常填写。','1. lijunhua、admin 从内部入口分别登录。\n2. lixiaolong、gysadmin 从门户入口分别登录。\n3. 在门户查看自己的企业名称；记下名称，后续邀请该企业。\n4. 查看你此前发布的受邀寻源是否可见；此时只查看，不提交测试报价。','内部：/?internal=1\n门户：/?portal=1\n密码用当前统一测试密码。','内部账号进入内部工作台；两个门户账号进入寻源列表，并可看到本企业收到的已发布询价。门户不能新建内部询价。')
case('M02',P0,'创建询报价草稿','lijunhua','已知道门户绑定的企业；该企业合格，品类与询价一致。','1. 打开“寻源与竞价”→“新建询价”。\n2. 策略选“询报价”，邀请M01记下的企业。\n3. 填下列测试数据并“保存草稿”。\n4. 记录系统生成的询价编号。\n5. 门户搜索该标题，此时应看不到。','标题：验收-询报价-001（复测换001后缀）\n编号留空自动生成；CNY\n物料：验收测试件；数量10\n交付要求：7天；基准单价110（可选）\n截止：当前时间后10分钟。','内部保存为“草稿”，数量/交期正确，刷新后仍存在；未发布草稿不在门户显示。若找不到企业，先核对品类和“合格”状态。')
case('M03',P0,'上传并保存物料规格','lijunhua','M02草稿已保存。','1. 打开草稿，“规格所属物料”选验收测试件。\n2. “上传规格”选择配套PDF。\n3. 上传后会进入编辑，务必再次点“保存草稿”。\n4. 再打开/编辑，检查绑定数量和原有数量、交期。','配套文件：W12-15-验收规格示例.pdf\n仅限本地验收；PDF/PNG/JPEG，≤10MB。','相应物料有已绑定规格；下载能正常打开；普通编辑不会丢失物料数据。多行时应绑定到所选行，而非总是第一行。')
case('M04',P0,'发布邀请与核心项锁定','lijunhua → lixiaolong/gysadmin','M02/M03草稿数据和规格已保存，仍未到截止。','1. lijunhua 点“发布邀请”。\n2. 门户刷新，搜索验收标题并“打开询价”。\n3. 核对数量、交期、规格并下载附件。\n4. 返回内部，检查发布后核心物料/邀请不再提供草稿编辑。','同M02单；受邀企业必须与门户绑定企业相同。','状态为“报价中”；两个门户身份都能看到同一份邀请和规格；发布后核心物料与邀请名单不可随意修改。')
case('M05',P0,'供应商第一轮报价','lixiaolong','M04已发布；当前时间仍在截止前。','1. 门户打开本次询价。\n2. 输入单价、交期、运费及说明。\n3. 点“提交新一轮报价”，等成功后重新打开检查。','数量由RFQ固定为10\n含税单价100；交期7天\n运费20；说明：第一轮验收。','第一轮报价留存；到厂总额应为 10×100+20=1020。单价、交期、运费与输入一致。')
case('M06',P0,'第二轮报价与历史保留','lixiaolong','第一轮已确认成功；仍在截止前。','1. 在同一询价改为第二轮数据。\n2. 再点“提交新一轮报价”。\n3. 刷新或重新打开，检查两轮记录。','含税单价95；交期7天；运费20\n说明：第二轮验收。','第二轮总额970；第一轮1020仍可查，不能被覆盖。最新报价应为第二轮，而不是累计两轮金额。')
case('M07',P0,'同企业两个门户身份共享业务','gysadmin','M06已有两轮报价。','1. 退出/切换到gysadmin登录门户。\n2. 打开同一询价，查看报价轮次和附件。\n3. 此项只查看，不再提交第三轮。','本次询价，第一轮1020、第二轮970。','能看到同企业的两轮报价和规格。两账号代表同一家企业，共享本方报价是正确行为，并非越权。')
case('M08',P0,'比价、实际绩效与Excel导出','lijunhua','M06至少有一轮有效报价；可在等待截止时完成。','1. 内部打开询价，点“查看最新比价”。\n2. 核对最新总额、交期、绩效有效月份和风险。\n3. 点“导出比价 Excel”，打开下载文件。','最新总额970；交期7天\n只有一家有效报价时，成本分应为100。','比价采用最新轮次970；导出包含相同企业及当前总额/评分信息。无绩效历史应明确显示无历史，综合分不伪造0分或满分；风险UNKNOWN表示未评估，不是无风险。')
case('M09',P0,'截止封盘并提交定点','lijunhua','等待页面显示的实际截止时间；必填资质完整且有效。','1. 到截止后稍等数秒，点“刷新寻源”并重新打开。\n2. 确认“已封盘”。\n3. 选择有有效报价的企业，填写推荐原因；无历史时填例外原因。\n4. 点“提交定点申请”，记录申请编号。','推荐原因：本地验收试单，报价已核对。\n无历史例外：仅用于验收试单，资质已核验，后续跟踪绩效。','未到截止不应能定点；封盘后可提交，显示“定点审批中/PENDING”。缺失/过期资质或没有有效报价应被明确阻断；补齐前置后再测。')
case('M10',P0,'禁止审核本人定点申请','lijunhua','M09已由lijunhua提交，尚未由admin审批。','1. 查看该定点申请的审批详情。\n2. 若显示审批操作，用lijunhua尝试“批准”并填意见。\n3. 检查待办/申请状态。','意见：本人申请权限验证。','自己的待办不出现本申请，或提交审批时明确拒绝本人自审，均为正确行为。申请保持待审批，不得被本人批准。')
case('M11',P0,'管理员审阅并批准定点','admin','M09申请待审批；操作与验收标题相同的待办，勿误批其他单。','1. 在“定点审批待办”打开本次申请。\n2. 核对供方、总额970、绩效月数、风险及理由。\n3. 如为无历史，明确勾选批准例外；HIGH/UNKNOWN风险需勾选确认风险。\n4. 动作选“批准”，填写意见后“处理定点审批”。','意见：验收数据已核对，同意试单。\n无历史/HIGH/UNKNOWN仅在审阅后勾选。','相应例外/风险未确认时应阻断；确认并提交后为APPROVED/定点待发布。注意：定点当前是经理审批；不是供应商准入的SQE+经理两级流程。')
case('M12',P0,'发布定点及门户本方结果','lijunhua → 两个门户账号','M11已通过；审批依据没有再改动。','1. lijunhua刷新询价、打开已批准的定点记录。\n2. 点“发布定点与落选通知”。\n3. 两个门户账号分别刷新并打开该询价，查看结果和本方通知。\n4. 记录单号和截图。','同M02主线单，不能拿其他已完成单反复审批。','内部为“已定点”；两个门户身份都能看到本企业中选结果/通知，不能查看内部完整审批快照。只有一家供方，本项不验证多供方落选通知。')
case('A01',P0,'新建快速竞价并首次出价','lijunhua → lixiaolong','用新单；仍邀请门户绑定企业。','1. 新建“反向竞价”，填下列参数并发布。\n2. lixiaolong门户打开，等“已连接”。\n3. 在“新的篮子总价”输入990并“发送竞价”。\n4. 等持久化成功回执。','标题：验收-竞价-001\n开始时间：当前分钟；截止：当前后3分钟\n起始整单总价1000；最小降幅10\n末段顺延分钟0（关闭，便于快速封盘）\n物料数量10；本次输入990是整单总价，不是单价。','连接成功；990被接受，最低总价和本方报价为990，本方排名1。匿名榜单用“供方-…”席位；页面顶部显示自己的企业名正常。')
case('A02',P0,'最小降幅拒绝与再次合法出价','lixiaolong','A01最新最低价990，最小降幅10，尚未截止。','1. 尝试输入985。\n2. 应明确拒绝，最低价仍990。\n3. 如页面保留待确认请求，可查原回执；明确未成交后点“确认未成交后重填”。\n4. 改填980并发送，等成功回执。','先985（降幅5，非法）\n再980（降幅10，合法）。','985不能改变最低价或排名；980成功后本方/最低价为980。不要在未确认时连续换新请求重复出价。')
case('A03',P0,'刷新重连、保存与截止封盘','lixiaolong → lijunhua','A02已收到980的成功回执；本单顺延为0。','1. 刷新浏览器，回到列表后重新打开同一竞价。\n2. 等“已连接”，核对最低价、本方价和排名。\n3. 等实际截止后数秒，再点“刷新寻源”重新打开。','同A01/A02单；已确认价980。','刷新后980和排名仍在，不自动重复提交。到点后为“已封盘”，不能再接受新价；可用M09–M12另验竞价定点。')
case('A04',P1,'防狙击延时与连续延时','lijunhua → lixiaolong','另建新竞价；不要修改A01已发布单的核心参数。','1. 新单起始1000、降幅10，截止设当前后3分钟，顺延设1分钟。\n2. 在原截止前最后1分钟内出990，记下新截止。\n3. 再等到新的最后1分钟内出980。\n4. 观察截止再次后移。','标题：验收-延时-001\n末段顺延1分钟\n建议在每次截止前30～50秒出价。','每笔在最后窗口内的合法出价，将原截止再加1分钟；第二笔也可继续延时。若第二笔立即提交、距新截止已超过1分钟，不触发延时属正常。')
case('A05',P1,'无人报价自动流标','lijunhua','另建新竞价，顺延0；不让任何门户账号出价。','1. 新建并发布，截止设当前后2～3分钟。\n2. 全程不出价。\n3. 到点后数秒刷新、重新打开。','标题：验收-流标-001\n起始1000；降幅10；顺延0。','状态为“流标”，没有虚构中选候选，也不能无报价直接发布定点。')
case('E01',P1,'未报价撤回与草稿保密','lijunhua → 两个门户账号','用两个新单，不影响主线。','1. 新建草稿“验收-撤回草稿”，不发布即撤回；门户搜标题并看通知。\n2. 另建“验收-未报价撤回”，发布但不报价，再撤回。\n3. 门户刷新查看。','两张新询价，截止均设未来10分钟。','未发布草稿即使撤回也不能泄露到门户/通知。已发布但无报价的单可以撤回，门户显示撤回状态/通知且不能报价。')
case('E02',P1,'已有报价不能撤回','lijunhua','主线M05之后、M09截止之前可测；或另建已报价单。','1. 打开已有报价的询价。\n2. 点“撤回未报价询价”。\n3. 核对提示与已有报价。','已有至少一轮报价；不得用未报价单替代。','明确拒绝：已有报价或出价，不能撤回。状态/截止及原报价不丢失。')
case('E03',P1,'延期与供应商通知','lijunhua → 两个门户账号','报价中且尚未截止；建议用新单，若用主线要按新截止等待。','1. 记录原截止。\n2. “延长至”选更晚2分钟，填原因后“延长截止并通知”。\n3. 两门户刷新，核对截止和本方通知。','原因：验收延期通知。\n新截止=原截止+2分钟。','截止确实后移，门户可见延期通知。不能缩短截止或选过去时间；不应清空已有报价。')
case('E04',P1,'驳回后重新提交定点','admin → lijunhua','用另一张已报价且封盘的测试单；或在主线M11批准前走此分支。','1. lijunhua提交定点，admin动作选“驳回”，填理由。\n2. lijunhua刷新，查看驳回意见和记录。\n3. 补充推荐说明后重新提交；再按M11/M12批准发布。','驳回：请补充试单和跟踪措施。\n重提：已补充试单数量及绩效跟踪安排。','旧申请保留REJECTED和意见；询价回到已封盘，可生成新的待审批申请，不覆盖旧记录。已发布定点不能拿来当此分支的待审批单。')
case('E05',P1,'两个门户身份的管理权限差异','lixiaolong / gysadmin','只检查界面和查询，不变更当前账号角色、不停用自己。','1. 两账号分别打开“供应商档案”，查看企业范围。\n2. 打开该企业的联系人/资质页，比较可维护操作。\n3. 检查门户没有内部新建供应商/新建询价能力。','lixiaolong：业务测试；gysadmin：本企业管理员。','两者都只属于同一企业；gysadmin可维护本企业资料/资质/联系人，lixiaolong不能管理联系人权限。合格档案的核心变更仍按审批规则，不能用直接改状态来替代审批。')
case('E06',P1,'原资质上传及两级准入回归（建议一次）','lijunhua发起；admin审批两次','另建“验收-准入-001”；不要覆盖原供应商资质。','1. lijunhua新建供应商，填说明页的合成测试信用代码。\n2. 新增启用的管理员联系人即可，不开第五个登录账号。\n3. 按“资质证书”缺失清单逐项录入，上传示例PDF、填写有效日期。\n4. 在“准入与审批”提交准入，lijunhua不得审核自己。\n5. admin从“审批工作台”依次完成SQE、经理节点。','名称：验收-准入-001（每次换后缀与新代码）\n电话13800138000；证书编号UAT-类型\n生效日今天或之前；到期日未来一年。\n示例PDF只用于本机新验收供方。','缺资质时准入应阻断；齐备后可提交。第一节点后仍审核中，第二节点后为合格，两个节点/意见有记录。现有门户账号仍只绑定原企业，不能自动看到新建企业。')
assert len(cases)==23
alphabet='0123456789ABCDEFGHJKLMNPQRTUWXY';weights=[1,3,9,27,19,26,16,17,20,29,25,13,8,24,10,30,28]
def credit():
    body='9111'+str(int(datetime.now().timestamp()*1000))[-10:]+str(secrets.randbelow(900)+100)
    return body+alphabet[(31-sum(alphabet.index(c)*w for c,w in zip(body,weights))%31)%31]
codes=[]
while len(codes)<3:
    code=credit()
    if code not in codes:codes.append(code)
notes=[
('本轮目标','先跑“主线必测”M01–M12，再跑竞价A01–A03。覆盖寻源→报价/竞价→定点审批→本方结果。主线约20～30分钟，含等待；选测另算。'),
('内部入口','http://localhost:18080/?internal=1'),('供应商入口','http://localhost:18080/?portal=1'),
('账号分工','lijunhua：内部测试/发布/发起\nadmin：管理及审批\nlixiaolong：供应商报价/业务测试\ngysadmin：本供应商资料/资质/联系人管理'),
('密码','沿用已设置的统一测试密码；不在本表写密码。当前账号文件：artifacts/SRM-local-review/runtime/登录账号.txt。'),
('先核对绑定企业','用门户账号查看企业名称。两个门户账号属于同一家受邀企业；新询价必须邀请这家企业，品类也要与其档案一致。管理员账号不代表其他所有供应商。'),
('窗口与账号切换','可用普通窗口登录内部、无痕窗口登录门户；换账号先退出。不要多处同时登录同一账号，新登录会使旧会话失效。手测时不要同时运行使用这四个账号的自动脚本。'),
('记录位置','在每张用例表的“关联单号、实际结果、测试结论、截图/备注”填写；结论支持下拉。开始前记下本次询价编号和实际截止。'),
('固定验收金额','询报价：数量10，第一轮100×10+20=1020；第二轮95×10+20=970。\n竞价：输入整单总价，起始1000、降幅10，990合法、985（在990后）非法、980合法。'),
('封盘与审批别混淆','询报价截止设当前后10分钟，不用默认明天。定点要等实际截止后自动封盘，再刷新重新打开。定点当前是经理审批；供应商准入才是SQE+经理两级。'),
('没有历史不是故障','无历史绩效应显示无历史/综合分空，不能冒填0或100；定点要写例外理由并由admin明确批准。UNKNOWN风险表示未评估，需明确审阅确认，不表示无风险。'),
('测试数据隔离','全部新单/新供应商冠以“验收-…”并写明测试用途。已有寻源仅做可见性查询，不向原单提交测试报价，不覆盖原证照，不清空数据库。'),
('结果判定','通过：达到预期，包括正确阻止非法操作。\n失败：前置满足但结果不符。\n阻塞：前置数据不足/环境问题暂时无法执行，填写原因。\n不适用：当前数据/账号确实无法覆盖，不冒填通过。'),
('当前四账号的覆盖边界','两门户身份是同一企业，看到同一报价/本方结果正常。当前不安排多供方落选、跨供方私聊、改派其他供方、经理成功转办或50方并发的人工测试。它们需要额外独立身份/数据，历史自动化证据另存。'),
('本次不验后续模块','真实外部风险源、ERP/MES/QMS集成、订单/ASN/对账等不在本次W12–15手工验收范围。未知风险按现有人工依据/确认规则处理。'),
('示例上传文件','同目录：W12-15-验收规格示例.pdf。可作为本机技术规格样例；E06只能给新建验收供方模拟资质上传，不是正式证书。'),
('E06合成信用代码1',codes[0]),('E06合成信用代码2',codes[1]),('E06合成信用代码3',codes[2]),
('信用代码说明','这些是带校验位的合成测试数据，不代表真实企业；仅用于新建本机验收记录。每次复测使用未使用过的代码，不复制既有供应商代码。'),
('主线询价编号/截止','待填写：'),('快速竞价编号/截止','待填写：'),('延时竞价编号/截止','待填写：'),('验收日期/人员','待填写：'),('总体结论','未开始；先完成15项P0，P1按需要选做。')]
wb=Workbook();intro=wb.active;intro.title='使用说明';intro.append(['项目','说明/可填写内容'])
for row in notes:intro.append(row)
intro.column_dimensions['A'].width=27;intro.column_dimensions['B'].width=112
main=wb.create_sheet('主线必测');extra=wb.create_sheet('竞价与选测')
widths=[15,12,30,24,48,88,62,92,25,55,16,42]
for sheet,data in [(main,[r for r in cases if r[0].startswith('M')]),(extra,[r for r in cases if not r[0].startswith('M')])]:
    sheet.append(headers)
    for row in data:sheet.append(row)
    for i,width in enumerate(widths,1):sheet.column_dimensions[get_column_letter(i)].width=width
    sheet.auto_filter.ref=sheet.dimensions
    dv=DataValidation(type='list',formula1='"未测,通过,失败,阻塞,不适用"',allow_blank=True);dv.errorTitle='请选择测试结论';dv.error='使用下拉选项：未测、通过、失败、阻塞、不适用';dv.showErrorMessage=True;sheet.add_data_validation(dv);dv.add(f'K2:K{sheet.max_row}')
    for value,color in [('通过','E8F4EC'),('失败','FCE8E6'),('阻塞','FFF4CE')]:sheet.conditional_formatting.add(f'A2:L{sheet.max_row}',FormulaRule(formula=[f'$K2="{value}"'],fill=PatternFill('solid',fgColor=color)))
issues=wb.create_sheet('问题记录');issues.append(['问题编号','关联用例','账号','关联单号/供方','复现步骤','实际现象','预期结果','截图文件/发生时间','严重程度','处理状态'])
for _ in range(20):issues.append(['']*10)
for i,width in enumerate([18,18,22,32,80,65,65,45,18,20],1):issues.column_dimensions[get_column_letter(i)].width=width
for col,values in [('I','阻断,严重,一般,建议'),('J','待确认,处理中,已修复,复测通过')]:
    dv=DataValidation(type='list',formula1='"'+values+'"',allow_blank=True);issues.add_data_validation(dv);dv.add(f'{col}2:{col}1000')
issues.auto_filter.ref='A1:J21'
for sheet in wb:
    sheet.freeze_panes='A2';sheet.sheet_view.showGridLines=False
    for cell in sheet[1]:cell.font=Font(name='Microsoft YaHei',bold=True,color='FFFFFF',size=11);cell.fill=PatternFill('solid',fgColor='176B54');cell.alignment=Alignment(vertical='center',wrap_text=True)
    sheet.row_dimensions[1].height=30
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font=Font(name='Microsoft YaHei',size=11);cell.alignment=Alignment(vertical='top',wrap_text=True);cell.border=Border(bottom=Side(style='hair',color='DCE6DF'))
            if isinstance(cell.value,str) and cell.value.startswith('http://'):
                cell.hyperlink=cell.value;cell.font=Font(name='Microsoft YaHei',size=11,color='0563C1',underline='single')
        sheet.row_dimensions[row[0].row].height=100 if sheet in (main,extra) else (75 if sheet==issues else 66)
    sheet.sheet_properties.pageSetUpPr.fitToPage=True;sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A3;sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0;sheet.print_title_rows='1:1'
intro.sheet_properties.tabColor='176B54';main.sheet_properties.tabColor='1B8E64';extra.sheet_properties.tabColor='D69B23';issues.sheet_properties.tabColor='C44A45'
wb.save(paths[0])
with paths[2].open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.writer(f,lineterminator="\n");w.writerow(headers);w.writerows(cases)
md=['# W12–15 手工验收用例','',f'生成日期：{datetime.now():%Y-%m-%d}。本文件只生成文档，不登录系统、不修改账号或数据库。','',
    '推荐打开同目录 Excel，按“主线必测”→“竞价与选测”操作，逐项填写实际结果和结论。现有账号密码不写入文档。','',
    '## 开始前','']
for title,text in notes[:16]:md.extend([f'**{title}**：{text.replace(chr(10),"；")}',''])
md.extend(['## 测试数据','',*(f'- 合成信用代码{i+1}：`{code}`（仅E06新建本机验收记录，不代表真实企业）' for i,code in enumerate(codes)),'','## 顺序总览','', '| 编号 | 优先级 | 场景 | 账号 |','| --- | --- | --- | --- |'])
for row in cases:md.append('| '+' | '.join(row[:4])+' |')
for row in cases:
    cid,priority,title,actor,pre,steps,data,expected,*_=row
    md.extend(['',f'## {cid} · {title}（{priority}）','',f'账号：{actor}。','',f'前置：{pre}。','',steps,'','输入：','',data.replace('\n','<br>\n'),'','预期：','',expected,'','- [ ] 通过　- [ ] 失败　- [ ] 阻塞　- [ ] 不适用','', '关联单号：________　实际结果：________　截图/问题：________'])
md.extend(['','## 反馈问题的最小信息','','用例编号 + 使用账号 + 业务单号 + 发生时间 + 实际提示/截图。预期被阻止的操作如果确实被阻止，应填“通过”；不要把未到截止、无绩效历史或已正确拦截的非法报价当作程序故障。',''])
paths[1].write_text('\n'.join(md),encoding='utf-8')
stream=b'BT /F1 18 Tf 60 735 Td (SRM Local Acceptance - Technical Specification) Tj /F1 11 Tf 0 -35 Td (FOR LOCAL TESTING ONLY - NOT A CERTIFICATE) Tj 0 -28 Td (Material: acceptance test component / steel) Tj 0 -24 Td (Quantity: 10 / Delivery requirement: 7 days) Tj 0 -24 Td (Tolerance example: 0.01 mm) Tj 0 -32 Td (Use only for a new clearly named local acceptance record.) Tj 0 -24 Td (Do not replace existing supplier qualification documents.) Tj ET\n'
objects=[b'<< /Type /Catalog /Pages 2 0 R >>',b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>',b'<< /Length '+str(len(stream)).encode()+b' >>\nstream\n'+stream+b'endstream',b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>']
pdf=bytearray(b'%PDF-1.4\n%\xe2\xe3\xcf\xd3\n');offsets=[0]
for i,obj in enumerate(objects,1):offsets.append(len(pdf));pdf.extend(f'{i} 0 obj\n'.encode()+obj+b'\nendobj\n')
xref=len(pdf);pdf.extend(b'xref\n0 6\n0000000000 65535 f \n')
for offset in offsets[1:]:pdf.extend(f'{offset:010d} 00000 n \n'.encode())
pdf.extend(f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode());paths[3].write_bytes(pdf)
# Artifact integrity, not a claim that the user has executed the cases.
check=load_workbook(paths[0]);assert check.sheetnames==['使用说明','主线必测','竞价与选测','问题记录'];assert check['主线必测'].max_row==13 and check['竞价与选测'].max_row==12
assert len(check['主线必测'].data_validations.dataValidation)==1
with zipfile.ZipFile(paths[0]) as z:
    assert z.testzip() is None
    for name in z.namelist():
        if name.endswith('.xml'):ET.fromstring(z.read(name))
with paths[2].open(encoding='utf-8-sig',newline='') as f:assert len(list(csv.reader(f)))==24
assert all(len(c)==18 and sum(alphabet.index(ch)*w for ch,w in zip(c[:17],weights))%31==(31-alphabet.index(c[-1]))%31 for c in codes)
print(f'Generated and validated workbook, Markdown, CSV and sample PDF: {len(cases)} cases; 15 required, 8 optional. All execution results remain untested.')
print(out)
